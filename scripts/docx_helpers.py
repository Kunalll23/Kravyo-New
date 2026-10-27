import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_borders(cell, border_color="D3D3D3", border_sz="4", border_val="single"):
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="{border_val}" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
        <w:left w:val="{border_val}" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="{border_val}" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
        <w:right w:val="{border_val}" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))

def apply_table_styles(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            # Repeat header on subsequent pages
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        for c_idx, cell in enumerate(row.cells):
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            if r_idx == 0:
                set_cell_shading(cell, "E9ECEF")
                set_cell_borders(cell, border_color="6C757D", border_sz="6")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.name = 'Cambria'
                        r.font.size = Pt(10.5)
                        r.font.bold = True
            else:
                bg = "FFFFFF" if r_idx % 2 == 1 else "F8F9FA"
                set_cell_shading(cell, bg)
                set_cell_borders(cell, border_color="DEE2E6", border_sz="4")
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.name = 'Cambria'
                        r.font.size = Pt(10)

def add_image_placeholder(doc, caption, figure_id=None, hint="(File Reference / Output Image)"):
    # Check if hint points to an existing image file
    potential_paths = []
    if hint:
        raw_hint = hint.replace("File Reference: ", "").strip()
        potential_paths.append(os.path.join(r"c:\xampp\htdocs\Kravyo", raw_hint))
        potential_paths.append(os.path.join(r"c:\xampp\htdocs\Kravyo\diagrams", os.path.basename(raw_hint)))
        potential_paths.append(raw_hint)
    
    found_img = None
    for path in potential_paths:
        if os.path.exists(path) and os.path.isfile(path) and path.lower().endswith(('.png', '.jpg', '.jpeg')):
            found_img = path
            break
            
    if found_img:
        # Embed actual picture
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        r_img = p_img.add_run()
        r_img.add_picture(found_img, width=Inches(6.2))
    else:
        # 1x1 table placeholder
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.27)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        set_cell_shading(cell, "F8F9FA")
        set_cell_borders(cell, border_color="A0AAB2", border_sz="8", border_val="dashed")
        set_cell_margins(cell, top=300, bottom=300, left=200, right=200)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(28)
        p.paragraph_format.space_after = Pt(28)
        
        r1 = p.add_run(f"[ INSERT IMAGE HERE: {caption.upper()} ]\n")
        r1.font.name = 'Cambria'
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(0, 0, 0)
        
        r2 = p.add_run(f"{hint}\n")
        r2.font.name = 'Cambria'
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(0, 0, 0)
        
        r3 = p.add_run("Recommended Size: 1920 x 1080 px or 1280 x 720 px")
        r3.font.name = 'Cambria'
        r3.font.size = Pt(8.5)
        r3.font.italic = True
        r3.font.color.rgb = RGBColor(0, 0, 0)
        
    # Caption below
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.line_spacing = 1.5
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(6)
    
    r_cap_label = p_cap.add_run(f"Figure {figure_id}: " if figure_id else "Figure: ")
    r_cap_label.font.name = 'Cambria'
    r_cap_label.font.size = Pt(11)
    r_cap_label.font.bold = True
    r_cap_label.font.color.rgb = RGBColor(0, 0, 0)
    
    r_cap_text = p_cap.add_run(caption)
    r_cap_text.font.name = 'Cambria'
    r_cap_text.font.size = Pt(11)
    r_cap_text.font.italic = True
    r_cap_text.font.color.rgb = RGBColor(0, 0, 0)

def add_table_caption(doc, caption, table_id=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    if table_id:
        r_lbl = p.add_run(f"Table {table_id}: ")
        r_lbl.font.name = 'Cambria'
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(0, 0, 0)
    r_txt = p.add_run(caption)
    r_txt.font.name = 'Cambria'
    r_txt.font.size = Pt(11)
    r_txt.font.bold = True
    r_txt.font.italic = True
    r_txt.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_chapter_title(doc, title_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(title_text)
    run.font.name = 'Cambria'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_main_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Cambria'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_sub_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Cambria'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_sub_sub_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = 'Cambria'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Cambria'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    run = p.add_run(text)
    run.font.name = 'Cambria'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Cambria'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    run = p.add_run(text)
    run.font.name = 'Cambria'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def setup_header_footer(section, is_roman=False):
    # Setup page size A4 & 1 inch margins
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Set page numbering format
    sectPr = section._sectPr
    pgNum = sectPr.find(qn('w:pgNumType'))
    if pgNum is None:
        pgNum = OxmlElement('w:pgNumType')
        sectPr.append(pgNum)
    pgNum.set(qn('w:fmt'), 'lowerRoman' if is_roman else 'decimal')
    if not is_roman:
        pgNum.set(qn('w:start'), '1')
        
    # Header: Project title in solid black (Bold Italic Cambria)
    header = section.header
    header.is_linked_to_previous = False
    p_head = header.paragraphs[0]
    p_head.text = ""
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_head.paragraph_format.space_before = Pt(0)
    p_head.paragraph_format.space_after = Pt(0)
    r_head = p_head.add_run("Kravyo: A Multi-Tenant Cloud Kitchen & Meal Subscription Platform")
    r_head.font.name = 'Cambria'
    r_head.font.size = Pt(9.5)
    r_head.font.italic = True
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(0, 0, 0)
    
    # Footer: 1-row, 2-column borderless table in solid black
    footer = section.footer
    footer.is_linked_to_previous = False
    # Clear any inherited tables in footer
    for t in list(footer.tables):
        t._element.getparent().remove(t._element)
        
    p_foot_init = footer.paragraphs[0]
    p_foot_init.text = ""
    
    tbl_foot = footer.add_table(rows=1, cols=2, width=Inches(6.27))
    tbl_foot.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left = tbl_foot.cell(0, 0)
    c_right = tbl_foot.cell(0, 1)
    c_left.width = Inches(4.7)
    c_right.width = Inches(1.57)
    
    # Remove borders from footer table
    set_cell_borders(c_left, border_color="FFFFFF", border_sz="0", border_val="none")
    set_cell_borders(c_right, border_color="FFFFFF", border_sz="0", border_val="none")
    set_cell_margins(c_left, top=0, bottom=0, left=0, right=0)
    set_cell_margins(c_right, top=0, bottom=0, left=0, right=0)
    
    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_l.paragraph_format.line_spacing = 1.0
    p_l.paragraph_format.space_before = Pt(0)
    p_l.paragraph_format.space_after = Pt(0)
    r_enr = p_l.add_run("202407100510063 | 202407100510077 | 202407100510079 | 202407100510121")
    r_enr.font.name = 'Cambria'
    r_enr.font.size = Pt(9)
    r_enr.font.bold = True
    r_enr.font.color.rgb = RGBColor(0, 0, 0)
    
    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.line_spacing = 1.0
    p_r.paragraph_format.space_before = Pt(0)
    p_r.paragraph_format.space_after = Pt(0)
    r_pg_label = p_r.add_run("Page ")
    r_pg_label.font.name = 'Cambria'
    r_pg_label.font.size = Pt(9)
    r_pg_label.font.bold = True
    r_pg_label.font.color.rgb = RGBColor(0, 0, 0)
    
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Cambria" w:hAnsi="Cambria"/><w:b/><w:sz w:val="18"/><w:color w:val="000000"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>' % nsdecls('w'))
    p_r._p.append(fldSimple)

print("Base helper functions initialized.")
