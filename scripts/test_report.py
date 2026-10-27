import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Set Normal Style
style = doc.styles['Normal']
style.font.name = 'Cambria'
style.font.size = Pt(12)
style.font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after = Pt(0)

# Section setup
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)

# Header
header = section.header
p_head = header.paragraphs[0]
p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_head.paragraph_format.space_before = Pt(0)
p_head.paragraph_format.space_after = Pt(0)
r_head = p_head.add_run("Kravyo: A Multi-Tenant Cloud Kitchen & Meal Subscription Platform")
r_head.font.name = 'Cambria'
r_head.font.size = Pt(9)
r_head.font.italic = True
r_head.font.color.rgb = RGBColor(100, 100, 100)

# Footer
footer = section.footer
p_foot = footer.paragraphs[0]
p_foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_foot.paragraph_format.space_before = Pt(0)
p_foot.paragraph_format.space_after = Pt(0)
r_foot_left = p_foot.add_run("202407100510063 | 202407100510077 | 202407100510079 | 202407100510121\t\tPage ")
r_foot_left.font.name = 'Cambria'
r_foot_left.font.size = Pt(9)
r_foot_left.font.color.rgb = RGBColor(100, 100, 100)

# Add page number XML
fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
p_foot._p.append(fldSimple)

doc.save('test_styled.docx')
print('test_styled.docx generated successfully!')
