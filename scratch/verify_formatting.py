import docx
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

doc = docx.Document('Kravyo_Project_Report.docx')

print("=== CHECKING FORMATTING GUIDELINES ===")

# 1. Page size & margins
sec_ok = True
for idx, sec in enumerate(doc.sections):
    w_ok = abs(sec.page_width.inches - 8.27) < 0.05
    h_ok = abs(sec.page_height.inches - 11.69) < 0.05
    m_top = abs(sec.top_margin.inches - 1.0) < 0.01
    m_bot = abs(sec.bottom_margin.inches - 1.0) < 0.01
    m_left = abs(sec.left_margin.inches - 1.0) < 0.01
    m_right = abs(sec.right_margin.inches - 1.0) < 0.01
    print(f"Section {idx}: A4 Size={w_ok and h_ok} ({sec.page_width.inches:.2f}x{sec.page_height.inches:.2f}), 1 inch Margins={m_top and m_bot and m_left and m_right}")
    if not (w_ok and h_ok and m_top and m_bot and m_left and m_right):
        sec_ok = False

# 2. Header and Footer
print("\n--- Header & Footer ---")
for idx, sec in enumerate(doc.sections):
    hdr_text = " ".join([p.text for p in sec.header.paragraphs]).strip()
    ftr_text = " ".join([c.text for t in sec.footer.tables for r in t.rows for c in r.cells] + [p.text for p in sec.footer.paragraphs]).strip()
    print(f"Section {idx} Header: {hdr_text}")
    print(f"Section {idx} Footer: {ftr_text}")
    has_proj_title = "kravyo" in hdr_text.lower()
    has_enr = "202407100510063" in ftr_text
    has_pg = "page" in ftr_text.lower()
    print(f"  -> Has Project Title: {has_proj_title}, Has Enrollment Nos: {has_enr}, Has Page Ref: {has_pg}")

# 3. Numbering (Roman before Chapter 1, Decimal after)
print("\n--- Page Numbering Type ---")
for idx, sec in enumerate(doc.sections):
    pg = sec._sectPr.find(qn('w:pgNumType'))
    fmt = pg.get(qn('w:fmt')) if pg is not None else "None"
    start = pg.get(qn('w:start')) if pg is not None else "None"
    print(f"Section {idx}: fmt={fmt}, start={start}")

# 4. Heading styles & Fonts
print("\n--- Heading Styles Check ---")
font_violations = 0
for p in doc.paragraphs:
    for r in p.runs:
        if r.font.name and r.font.name != "Cambria":
            font_violations += 1

print(f"Non-Cambria font runs in paragraphs: {font_violations}")

# Check specific headings:
headings_to_check = [
    ("Chapter 1: Project Requirement Document Submission (Sequence No. 3)", 20, True, False),
    ("1.1 Purpose", 16, True, False),
    ("1.2.1 Multi-Tenant Role Isolation & Authentication", 12, True, False),
    ("References", 20, True, False),
]

for title_query, exp_sz, exp_bold, exp_italic in headings_to_check:
    matched = False
    for p in doc.paragraphs:
        if title_query in p.text and p.runs:
            r = p.runs[0]
            sz = r.font.size.pt if r.font.size else "Default"
            b = bool(r.font.bold)
            it = bool(r.font.italic)
            ls = p.paragraph_format.line_spacing
            sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
            sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else 0
            print(f"Heading '{title_query[:35]}...': Size={sz}pt (exp {exp_sz}), Bold={b} (exp {exp_bold}), LineSpacing={ls}, Sb={sb}, Sa={sa}")
            matched = True
            break
    if not matched:
        print(f"Heading '{title_query}' NOT FOUND in paragraphs")
